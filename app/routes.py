from flask import Blueprint, render_template, request, jsonify, send_file
from io import BytesIO
from datetime import datetime, date, time
from pathlib import Path
import base64, re, json, numbers
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

bp = Blueprint('main', __name__)

# ------------------------------------------------------------------ schema
# shared/schema.json is the single source of truth for the proforma layout.
# The browser build (docs/) compiles the same file into docs/schema.js via
# build_static.py, so the two implementations cannot drift apart.
SCHEMA_PATH = Path(__file__).resolve().parent.parent / 'shared' / 'schema.json'
SCHEMA = json.loads(SCHEMA_PATH.read_text(encoding='utf-8'))

FIELDS = SCHEMA['fields']
SEARCH_FIELDS = SCHEMA['search_fields']
REQUIRED_FIELDS = SCHEMA['required_fields']
FORM_ORDER = SCHEMA['form_order']
WIDE_FIELDS = set(SCHEMA['wide_fields'])
LABELS = SCHEMA['labels']
TEMPLATE_HEADERS = SCHEMA['template_headers']
PROPERTY_TYPES = SCHEMA['property_types']
_FLOAT_FIELDS = set(SCHEMA['float_fields'])
HIGHLIGHT_ARGB = SCHEMA['highlight_argb']
MAX_RESULTS = SCHEMA['max_results']
STICKY_FIELDS = SCHEMA['sticky_fields']
OPTION_FIELDS = SCHEMA['option_fields']
CHOICE_FIELDS = SCHEMA['choice_fields']
PREFIX_FIELDS = SCHEMA['prefix_fields']

# Config the browser needs; injected into the page by index() / build_static.py.
UI_CONFIG = {k: SCHEMA[k] for k in (
    'search_fields', 'sticky_fields', 'option_fields', 'choice_fields',
    'prefix_fields', 'max_results')}

XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
XLSM_MIME = 'application/vnd.ms-excel.sheet.macroEnabled.12'


def norm(v):
    """Lower-case and strip every non-alphanumeric character."""
    if v is None:
        return ''
    return re.sub(r'[^a-z0-9]+', '', str(v).strip().lower())


# Header lookups go through norm(), which strips every non-alphanumeric
# character, so the raw keys are normalised the same way here. (Before this,
# raw keys like 'parcel number' / 'p/s' could never match a normalised header
# 'parcelnumber' / 'ps', so Parcel, Name, CNIC and House-code were never
# detected and every search came back empty.)
HEADER_MAP = {norm(k): v for k, v in SCHEMA['header_map'].items()}


def cell_text(v):
    """Excel cell -> plain JSON-safe string (dates and floats included)."""
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, datetime):
        return v.isoformat(sep=' ')
    if isinstance(v, (date, time)):
        return v.isoformat()
    if isinstance(v, numbers.Number):
        f = float(v)
        if f.is_integer() and abs(f) < 1e15:
            return str(int(f))
        return format(f, '.10g')
    return str(v).strip()


# ------------------------------------------------------------- workbook I/O
def wb_from_b64(s):
    # keep_vba keeps a macro-enabled (.xlsm) workbook intact on re-save.
    return load_workbook(BytesIO(base64.b64decode(s)), data_only=False, keep_vba=True)


def b64_from_wb(wb):
    bio = BytesIO()
    wb.save(bio)
    return base64.b64encode(bio.getvalue()).decode()


def is_macro(wb):
    return getattr(wb, 'vba_archive', None) is not None


def download_name(wb):
    return 'Parcel_Mapping_Updated.xlsm' if is_macro(wb) else 'Parcel_Mapping_Updated.xlsx'


# --------------------------------------------------------- header detection
def _map_row(ws, r):
    """Map one sheet row to {field: column index}; first occurrence wins."""
    cols = {}
    for c in range(1, (ws.max_column or 1) + 1):
        field = HEADER_MAP.get(norm(ws.cell(r, c).value))
        if field and field not in cols:
            cols[field] = c
    return cols


def find_header(ws):
    """Locate the header row and its column map, with a best-effort fallback."""
    best_score, best_row, best_cols = 0, None, {}
    for r in range(1, min(ws.max_row or 1, 40) + 1):
        cols = _map_row(ws, r)
        if 'parcel' in cols and set(cols) & {'name', 'cnic', 'house_code'}:
            return r, cols
        score = len(cols) + (5 if 'parcel' in cols else 0)
        if score > best_score:
            best_score, best_row, best_cols = score, r, cols
    if best_row:
        return best_row, best_cols
    return 4, _map_row(ws, 4)


def last_data_row(ws, hr, cols):
    """Last row that actually holds data (ws.max_row also counts styled blanks)."""
    check = [cols[f] for f in SEARCH_FIELDS if f in cols] or list(cols.values())
    if not check:
        return hr
    for r in range((ws.max_row or hr), hr, -1):
        if any(ws.cell(r, c).value not in (None, '') for c in check):
            return r
    return hr


# -------------------------------------------------------------- record I/O
def blank_record():
    return {f: '' for f in FIELDS}


def row_to_record(ws, r, cols):
    rec = blank_record()
    for f in FIELDS:
        c = cols.get(f)
        if c:
            rec[f] = cell_text(ws.cell(r, c).value)
    return rec


def read_records(ws, hr, cols):
    """Every non-empty data row as {'row': n, ...fields}."""
    out = []
    for r in range(hr + 1, last_data_row(ws, hr, cols) + 1):
        rec = row_to_record(ws, r, cols)
        if not any(rec[f] for f in SEARCH_FIELDS):
            continue
        rec['row'] = r
        out.append(rec)
    return out


def search_records(records, q, limit=MAX_RESULTS):
    """Rank exact match > prefix match > contains, across P/S, CNIC, H/S and Name."""
    qn = norm(q)
    if not qn:
        return []
    scored = []
    for rec in records:
        best = None
        for f in SEARCH_FIELDS:
            v = norm(rec.get(f, ''))
            if not v:
                continue
            if v == qn:
                s = 0
            elif v.startswith(qn):
                s = 1
            elif qn in v:
                s = 2
            else:
                continue
            best = s if best is None else min(best, s)
        if best is not None:
            scored.append((best, rec.get('row', 0), rec))
    scored.sort(key=lambda t: (t[0], t[1]))
    return [r for _, _, r in scored[:limit]]


_PLACEHOLDERS = {norm(v) for v in SCHEMA['placeholder_values']}


def field_options(records, limit=60):
    """Distinct values already used in the file, most common first.

    The enumerator's own details repeat on every row, so offering the values
    the file already contains saves retyping them. 'Nill' and friends are
    filler written by apply_defaults, never real choices.
    """
    out = {}
    for f in OPTION_FIELDS:
        counts = {}
        for r in records:
            v = str(r.get(f, '') or '').strip()
            if not v or norm(v) in _PLACEHOLDERS:
                continue
            counts[v] = counts.get(v, 0) + 1
        out[f] = [v for v, _ in sorted(counts.items(),
                                       key=lambda kv: (-kv[1], kv[0].lower()))][:limit]
    return out


def common_prefix(values):
    """Longest run of leading digits shared by every numeric value given.

    House codes within one file share a prefix and differ only in the last
    digits, so the shared part can be pre-filled. Only all-digit values count,
    and at least two distinct ones are needed before a prefix means anything.
    """
    vals = sorted({s for s in (str(v or '').strip() for v in values) if s.isdigit()})
    if len(vals) < 2:
        return ''
    first, last = vals[0], vals[-1]          # LCP of a sorted set == LCP(first, last)
    i = 0
    while i < min(len(first), len(last)) and first[i] == last[i]:
        i += 1
    return first[:i]


def field_prefixes(records):
    return {f: common_prefix([r.get(f, '') for r in records]) for f in PREFIX_FIELDS}


def apply_defaults(rec):
    """Blank optional fields -> 'Nill'; blank Property Type -> 'Other'."""
    rec = {f: ('' if rec.get(f) is None else str(rec.get(f, '')).strip()) for f in FIELDS}
    if not rec['parcel'] or not rec['name']:
        return rec
    for f in FIELDS:
        if f not in ('parcel', 'name', 'property_type') and not rec[f]:
            rec[f] = 'Nill'
    if not rec['property_type']:
        rec['property_type'] = 'Other'
    return rec


def coerce(field, text):
    """Write numbers as numbers, but keep CNIC and house codes as text."""
    s = '' if text is None else str(text).strip()
    if s == '':
        return None
    if field in _FLOAT_FIELDS:
        try:
            return float(s)
        except ValueError:
            return s
    if field in ('parcel', 'sr') and re.fullmatch(r'[1-9]\d{0,8}', s):
        return int(s)
    return s


def ensure_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws['A1'] = 'Proforma for parcel mapping'
    ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = 'District name:'
    ws['C2'] = 'District ID:'
    ws['E2'] = 'Tehsil name:'
    ws['G2'] = 'Tehsil ID:'
    ws['I2'] = 'UC Name:'
    ws['L2'] = 'UC ID:'
    for i, h in enumerate(TEMPLATE_HEADERS, 1):
        c = ws.cell(4, i)
        c.value = h
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = 32 if i == 11 else 16
    ws.freeze_panes = 'A5'
    return wb


def load_state(b64):
    """Open a workbook (or a fresh template) and return (wb, ws, header_row, cols)."""
    wb = wb_from_b64(b64) if b64 else ensure_template()
    ws = wb[wb.sheetnames[0]]
    hr, cols = find_header(ws)
    return wb, ws, hr, cols


def sheet_payload(wb, ws, hr, cols):
    records = read_records(ws, hr, cols)
    return {
        'records': records,
        'options': field_options(records),
        'prefixes': field_prefixes(records),
        'header_row': hr,
        'sheet': ws.title,
        'columns': {f: get_column_letter(c) for f, c in cols.items()},
        'missing': [LABELS[f] for f in SEARCH_FIELDS if f not in cols],
        'macro': is_macro(wb),
        'downloadName': download_name(wb),
    }


# ------------------------------------------------------------------ routes
@bp.get('/')
def index():
    return render_template('index.html', fields=FORM_ORDER, labels=LABELS,
                           required=REQUIRED_FIELDS, wide=WIDE_FIELDS,
                           property_types=PROPERTY_TYPES, choices=CHOICE_FIELDS,
                           ui_config=UI_CONFIG)


@bp.post('/api/load')
def load():
    """Parse the uploaded workbook once; the browser then searches locally."""
    data = request.get_json(silent=True) or {}
    b64 = data.get('workbook', '')
    if not b64:
        return jsonify({'error': 'No workbook supplied.'}), 400
    try:
        wb, ws, hr, cols = load_state(b64)
    except Exception as e:
        return jsonify({'error': 'Could not read that file: %s' % e}), 400
    if 'parcel' not in cols:
        return jsonify({'error': 'No Parcel / P-S column found in this sheet. '
                                 'Check that the proforma header row is present.'}), 400
    payload = sheet_payload(wb, ws, hr, cols)
    payload['message'] = '%d record(s) loaded' % len(payload['records'])
    return jsonify(payload)


@bp.post('/api/search')
def search():
    """Server-side search - a fallback for clients without a local cache."""
    data = request.get_json(silent=True) or {}
    b64 = data.get('workbook', '')
    q = data.get('query', '')
    if not b64:
        return jsonify({'matches': [], 'message': 'No workbook loaded.'})
    try:
        wb, ws, hr, cols = load_state(b64)
        matches = search_records(read_records(ws, hr, cols), q)
    except Exception as e:
        return jsonify({'error': str(e)}), 400
    return jsonify({
        'matches': [{'row': m['row'], 'record': {f: m[f] for f in FIELDS}} for m in matches],
        'header_row': hr,
        'message': '%d match(es)' % len(matches),
    })


@bp.post('/api/save')
def save():
    data = request.get_json(silent=True) or {}
    b64 = data.get('workbook', '')
    rec = apply_defaults(data.get('record') or {})
    row = data.get('row')
    updating = row is not None

    missing = [LABELS[f] for f in REQUIRED_FIELDS if not rec[f]]
    if missing:
        verb = 'are' if len(missing) > 1 else 'is'
        return jsonify({'error': '%s %s required.' % (' and '.join(missing), verb)}), 400

    try:
        wb, ws, hr, cols = load_state(b64)
        if 'parcel' not in cols:
            return jsonify({'error': 'No Parcel / P-S column found in this sheet.'}), 400

        if updating:
            try:
                row = int(row)
            except (TypeError, ValueError):
                return jsonify({'error': 'Invalid row reference.'}), 400
            if row <= hr:
                return jsonify({'error': 'Invalid row reference.'}), 400
        else:
            existing = next((r for r in read_records(ws, hr, cols)
                             if norm(r['parcel']) == norm(rec['parcel'])), None)
            if existing:
                return jsonify({
                    'error': 'P/S %s already exists at row %d. Open it to update '
                             'instead of adding a duplicate.' % (rec['parcel'], existing['row']),
                    'duplicate': True,
                    'row': existing['row'],
                    'record': {f: existing[f] for f in FIELDS},
                }), 409
            row = last_data_row(ws, hr, cols) + 1
            sr_col = cols.get('sr')
            if sr_col and not ws.cell(row, sr_col).value:
                ws.cell(row, sr_col).value = row - hr

        for f in FIELDS:
            c = cols.get(f)
            if c:
                ws.cell(row, c).value = coerce(f, rec[f])

        yellow = PatternFill('solid', fgColor=HIGHLIGHT_ARGB)
        for c in range(1, max(cols.values()) + 1):
            ws.cell(row, c).fill = yellow

        out = b64_from_wb(wb)
    except Exception as e:
        return jsonify({'error': str(e)}), 400

    saved = {f: rec[f] for f in FIELDS}
    saved['row'] = row
    return jsonify({'workbook': out, 'record': saved, 'row': row,
                    'downloadName': download_name(wb), 'macro': is_macro(wb),
                    'action': 'updated' if updating else 'added'})


@bp.post('/api/download')
def download():
    data = request.get_json(silent=True) or {}
    b64 = data.get('workbook', '')
    if not b64:
        return jsonify({'error': 'Nothing to download yet - load a file or save a record first.'}), 400
    try:
        raw = base64.b64decode(b64)
    except Exception:
        return jsonify({'error': 'Workbook data is corrupt.'}), 400
    macro = bool(data.get('macro'))
    bio = BytesIO(raw)
    bio.seek(0)
    return send_file(
        bio, as_attachment=True,
        download_name='Parcel_Mapping_Updated.xlsm' if macro else 'Parcel_Mapping_Updated.xlsx',
        mimetype=XLSM_MIME if macro else XLSX_MIME)
