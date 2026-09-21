"""Turn a real parcel workbook into one that is safe to publish.

The live demo is served from a PUBLIC GitHub Pages site, so it must not carry
real people's details. This keeps everything that makes the demo useful - the
proforma layout, the header row, the column set, the row count, the parcel and
house-code numbering, the property types - and replaces everything that
identifies a person:

    Name        -> a made-up name, stable per row
    CNIC        -> a made-up but correctly shaped number
    First name  -> a made-up enumerator name
    Username    -> derived from that name
    Latitude    -> rounded to ~1 km, so it points at a district, not a house
    Longitude   -> likewise
    Comment     -> kept only when it is Add / Delete / Nill, else blanked

Formatting, column widths and the header row are untouched, because the point
of the demo is that it behaves exactly like the real thing.

    python tools/make_demo_workbook.py "path/to/real.xlsx" docs/sample-parcel-mapping.xlsx
"""
import sys
import hashlib
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook                      # noqa: E402
from app.routes import find_header, norm                # noqa: E402

FIRST_NAMES = ['Ahmad', 'Bilal', 'Usman', 'Hamza', 'Zain', 'Kashif', 'Naveed',
               'Imran', 'Tariq', 'Adnan', 'Faisal', 'Junaid', 'Rashid', 'Sajid',
               'Aisha', 'Fatima', 'Maryam', 'Sana', 'Hina', 'Nadia']
LAST_NAMES = ['Khan', 'Shah', 'Ahmed', 'Hussain', 'Iqbal', 'Malik', 'Aslam',
              'Rehman', 'Bibi', 'Nawaz', 'Javed', 'Siddiqui', 'Qureshi', 'Abbas']
ENUMERATORS = ['Ali', 'Sara', 'Naveed', 'Hina']
KEEP_COMMENTS = {'add', 'delete', 'nill'}


def seeded(value, salt):
    """A stable pseudo-random integer for a given cell, so reruns match."""
    h = hashlib.sha256(('%s|%s' % (salt, value)).encode('utf-8')).hexdigest()
    return int(h[:12], 16)


def fake_name(key, taken):
    """A made-up name that is not one of the real ones.

    Invented names can collide with real ones by chance - 'Fatima Bibi' did -
    and a real person's name appearing in the published demo is exactly what
    this script exists to prevent, coincidence or not. So keep re-rolling
    until the result is not in the source file.
    """
    for attempt in range(40):
        n = seeded('%s#%d' % (key, attempt), 'name')
        candidate = '%s %s' % (FIRST_NAMES[n % len(FIRST_NAMES)],
                               LAST_NAMES[(n // 97) % len(LAST_NAMES)])
        if norm(candidate) not in taken:
            return candidate
    return 'Resident %s' % (seeded(key, 'fallback') % 100000)


def strip_hidden_traces(wb, ws):
    """Clear the places a value survives after its cell has been overwritten.

    A filter dropdown remembers the exact values it was filtering on, so a
    CNIC can sit in the sheet XML long after the cell holding it was replaced.
    Document properties, notes and headers are the same kind of hiding place.
    """
    cleared = []
    try:
        from openpyxl.worksheet.filters import AutoFilter
        if ws.auto_filter is not None and (ws.auto_filter.ref or ws.auto_filter.filterColumn):
            cleared.append('auto-filter')
        ws.auto_filter = AutoFilter()
    except Exception:
        pass
    try:
        if ws.sort_state is not None:
            ws.sort_state = None
            cleared.append('sort state')
    except Exception:
        pass
    try:
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    cell.comment = None
                    cleared.append('cell note')
    except Exception:
        pass
    try:
        from openpyxl.packaging.core import DocumentProperties
        wb.properties = DocumentProperties()
        cleared.append('document properties')
    except Exception:
        pass
    try:
        for part in (ws.oddHeader, ws.oddFooter, ws.evenHeader,
                     ws.evenFooter, ws.firstHeader, ws.firstFooter):
            for side in (part.left, part.center, part.right):
                side.text = None
    except Exception:
        pass
    try:
        if len(wb.defined_names):
            wb.defined_names = type(wb.defined_names)()
            cleared.append('defined names')
    except Exception:
        pass
    return sorted(set(cleared))


def fake_cnic(key, like):
    """Same shape as the original: 5-7-1 digits, hyphenated or not."""
    n = seeded(key, 'cnic')
    a = 10000 + (n % 90000)
    b = 1000000 + ((n // 7) % 9000000)
    c = n % 10
    return ('%05d-%07d-%d' % (a, b, c)) if '-' in str(like) else ('%05d%07d%d' % (a, b, c))


SIGNIFICANT = 4      # 30.21 / 71.51 -> about 1 km


def coarse(value):
    """Blunt a coordinate to ~1 km so it no longer points at a house.

    Two encodings appear in these files: proper degrees (34.7712) and degrees
    with the point dropped (302126607 meaning 30.2126607, which is centimetre
    precision). Rounding only helps the first kind - the integer form survives
    untouched and still locates a doorstep - so the digits are cut in whichever
    form the cell uses.
    """
    try:
        f = float(value)
    except (TypeError, ValueError):
        return value
    if f == 0:
        return value
    if abs(f) < 1000:                       # ordinary degrees
        return round(f, 2)
    # Scaled integer: keep the leading significant digits, zero the rest.
    digits = '%d' % int(abs(f))
    if len(digits) <= SIGNIFICANT:
        return value
    blunted = int(digits[:SIGNIFICANT] + '0' * (len(digits) - SIGNIFICANT))
    return -blunted if f < 0 else blunted


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 1
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    if not src.exists():
        print('Source not found: %s' % src)
        return 1

    wb = load_workbook(src)
    ws = wb[wb.sheetnames[0]]
    hr, cols = find_header(ws)
    if 'parcel' not in cols:
        print('No Parcel/P-S column found; is this the right workbook?')
        return 1

    # Anything beyond the first sheet may hold notes about real people.
    for extra in wb.sheetnames[1:]:
        del wb[extra]

    # Every real name in the file, so an invented one can never match by chance.
    real_names = set()
    if 'name' in cols:
        for r in range(hr + 1, (ws.max_row or hr) + 1):
            v = ws.cell(r, cols['name']).value
            if v not in (None, ''):
                real_names.add(norm(v))

    changed = {'name': 0, 'cnic': 0, 'coords': 0, 'enumerator': 0, 'comment': 0}
    for r in range(hr + 1, (ws.max_row or hr) + 1):
        key = '%s:%s' % (r, ws.cell(r, cols['parcel']).value)

        for field, maker in (('name', lambda v: fake_name(key, real_names)),
                             ('cnic', lambda v: fake_cnic(key, v))):
            c = cols.get(field)
            if not c:
                continue
            cur = ws.cell(r, c).value
            if cur in (None, '') or norm(cur) in ('nill', 'na'):
                continue
            ws.cell(r, c).value = maker(cur)
            changed[field] += 1

        for field in ('latitude', 'longitude'):
            c = cols.get(field)
            if not c:
                continue
            before = ws.cell(r, c).value
            after = coarse(before)
            if after != before:
                ws.cell(r, c).value = after
                changed['coords'] += 1

        c = cols.get('first_name')
        if c and ws.cell(r, c).value not in (None, ''):
            who = ENUMERATORS[seeded(key, 'enum') % len(ENUMERATORS)]
            ws.cell(r, c).value = who
            changed['enumerator'] += 1
            u = cols.get('username')
            if u and ws.cell(r, u).value not in (None, ''):
                ws.cell(r, u).value = who.lower() + '01'

        c = cols.get('comment')
        if c:
            cur = ws.cell(r, c).value
            if cur not in (None, '') and norm(cur) not in KEEP_COMMENTS:
                ws.cell(r, c).value = 'Nill'
                changed['comment'] += 1

    cleared = strip_hidden_traces(wb, ws)

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)

    print('wrote %s' % dst)
    print('  hidden traces cleared: %s' % (', '.join(cleared) if cleared else 'none found')) 
    print('  header row kept at %d, columns kept: %s' % (hr, len(cols)))
    print('  names replaced      : %d' % changed['name'])
    print('  CNICs replaced      : %d' % changed['cnic'])
    print('  coordinates coarsened: %d' % changed['coords'])
    print('  enumerators replaced : %d' % changed['enumerator'])
    print('  comments blanked     : %d' % changed['comment'])
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
